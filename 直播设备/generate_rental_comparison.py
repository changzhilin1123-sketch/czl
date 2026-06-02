# -*- coding: utf-8 -*-
"""Generate live streaming equipment rental comparison Excel."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).resolve().parent / "直播租赁六家对比表_v2.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="微软雅黑", bold=True, size=14, color="1F4E79")
BODY_FONT = Font(name="微软雅黑", size=10)
NOTE_FONT = Font(name="微软雅黑", size=9, color="666666")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
ALT_FILL = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def style_header_row(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def style_data_area(ws, start_row, end_row, col_count, alt=False):
    for r in range(start_row, end_row + 1):
        for c in range(1, col_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if alt and (r - start_row) % 2 == 1:
                cell.fill = ALT_FILL


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_title(ws, title, note=None, merge_cols=12):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=merge_cols)
    row = 2
    if note:
        ws[f"A{row}"] = note
        ws[f"A{row}"].font = NOTE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=merge_cols)
        row += 1
    return row


def write_table(ws, start_row, headers, rows, col_widths):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h)
    style_header_row(ws, start_row, len(headers))
    for i, row_data in enumerate(rows):
        r = start_row + 1 + i
        for c, val in enumerate(row_data, 1):
            ws.cell(row=r, column=c, value=val)
    end_row = start_row + len(rows)
    style_data_area(ws, start_row + 1, end_row, len(headers), alt=True)
    set_col_widths(ws, col_widths)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return end_row + 2


def build_workbook():
    wb = Workbook()

    # ===== Sheet 1: 总览 =====
    ws1 = wb.active
    ws1.title = "总览"
    r = write_title(
        ws1,
        "直播租赁六家对比 — 总览（已更新折扣价）",
        "租期：3个月起 | 服务：六家均含上门调试 + 日常远程协助 | 场景：武隆天生三桥户外 + 室内绿幕",
        merge_cols=13,
    )
    headers1 = [
        "商家",
        "相机+灯光原价(元/月)",
        "相机+灯光折后(元/月)",
        "电脑(元/月)",
        "电脑配置摘要",
        "含电脑全套折后(元/月)",
        "3个月-拍摄端折后(元)",
        "3个月-含电脑折后(元)",
        "户外网络",
        "户外电源",
        "绿幕",
        "包邮/物流",
        "备注",
    ]
    rows1 = [
        [
            "租小杨", 2060, 1360, 300,
            "i5-12400F/32G/512G/RTX4060/27寸1K",
            1660, 4080, 4980,
            "无", "Power1000×1", "无", "需确认",
            "含电脑全套最低价；FX30+电源",
        ],
        [
            "昊天", 2110, 1470, 300,
            "i5-12400F/32G/512G/RTX3060 12G/27寸1K",
            1770, 4410, 5310,
            "无", "无", "无", "包邮寄出",
            "3个月起租；含顶光+魔术腿",
        ],
        [
            "阿水", "2080(折前)", 1456, 479,
            "i7 11代/32G/512G/RTX3060 6G 台式",
            3148.2, 4368, 9444.6,
            "聚合路由495", "Power1000×2", "无", "单程包邮(报价单)",
            "唯一含路由+双电源；全套最贵",
        ],
        [
            "猫冬", 2420, "1694(推算)", 303,
            "i5-12400F/32G/512G/RTX3060 6G 台式",
            1997, "5082(推算)", 5991,
            "无", "无", "无", "需确认",
            "折后全套1997/月；A7M4+300B×4",
        ],
        [
            "洋洋", 2420, 1694, 303,
            "i5 12代/32G/512G/RTX3060/27寸显示器 台式",
            1997, 5082, 5991,
            "无", "无", "无", "需确认",
            "A7M4+4灯；印迹200D+300B",
        ],
        [
            "易希", "2420.4(折前)", 1694.2, "不含/另租303",
            "可参考猫冬/洋洋同配置台式",
            "约1997(1694+303)", 5082.6, "约5991",
            "无", "无", "赠送", "单程包邮(报价单)",
            "唯一送绿幕；电脑需另租",
        ],
    ]
    write_table(
        ws1, r, headers1, rows1,
        [10, 14, 14, 10, 28, 16, 16, 16, 12, 14, 8, 12, 26],
    )

    # ===== Sheet 2: 相机采集 =====
    ws2 = wb.create_sheet("相机与采集")
    r = write_title(ws2, "相机与采集配置对比")
    headers2 = [
        "商家", "拍摄端原价(元/月)", "拍摄端折后(元/月)", "机身", "镜头", "UV镜",
        "采集卡", "HDMI线", "假电池", "备用电池+充电器", "64G卡+读卡器",
        "散热器", "三脚架等", "麦克风",
    ]
    rows2 = [
        ["租小杨", 2060, 1360, "Sony FX30", "索尼18-105", "无", "数魅VC20", "专用线", "有", "有", "有", "无", "三脚架+竖拍板+包", "猛玛一拖二"],
        ["昊天", 2110, 1470, "Sony A7C", "腾龙28-75", "无", "数魅VC20", "2条", "有", "有", "有", "无", "三脚架+竖拍板+包", "猛玛一拖二"],
        ["阿水", 2080, 1456, "Sony FX30", "适马18-50 F2.8", "无", "圆刚553", "5m×1", "有", "无", "无", "有", "三脚架+竖拍板+包", "猛玛M1+充电盒"],
        ["猫冬", 2420, 1694, "Sony A7M4", "腾龙28-75/适马28-70", "有", "数魅VC20 Pro", "绿联5m", "有", "有", "无", "有", "三脚架+竖拍板+包", "一拖二(品牌未写)"],
        ["洋洋", 2420, 1694, "Sony A7M4", "腾龙28-75/适马28-70", "有", "数魅VC20 Pro", "绿联5m", "有", "有", "有", "有", "三脚架+竖拍板+包", "一拖二(品牌未写)"],
        ["易希", 2420.4, 1694.2, "Sony A7M4", "适马28-70 F2.8", "无", "圆刚553", "5m×1", "有", "无", "无", "有", "三脚架+竖拍板+包", "猛玛M1+充电盒"],
    ]
    write_table(ws2, r, headers2, rows2, [10, 14, 14, 14, 18, 8, 14, 10, 8, 14, 14, 8, 16, 16])

    # ===== Sheet 3: 灯光配件 =====
    ws3 = wb.create_sheet("灯光与配件")
    r = write_title(ws3, "灯光与场景配件对比")
    headers3 = ["商家", "灯光(含于拍摄端)", "灯具", "单灯功率", "灯位设计", "绿幕", "户外电源", "户外网络"]
    rows3 = [
        ["租小杨", "含在1360内", "捷宝/凛光/数魅200W", "200W", "数量未详列", "无", "Power1000×1", "无"],
        ["昊天", "含在1470内", "捷宝200W×4(2单色+2双色)", "200W", "多柔光箱+顶光+魔术腿", "无", "无", "无"],
        ["阿水", "含在1456内", "KK 360WBi×4", "360W", "辅光/底子/造型/辅光", "无", "Power1000×2", "多网聚合路由"],
        ["猫冬", "含在1694内", "数魅300B×4", "300W", "2面光+2环境光", "无", "无", "无"],
        ["洋洋", "含在1694内", "印迹200D×2+数魅300B×2", "200W/300W", "2面光+2环境光(4灯)", "无", "无", "无"],
        ["易希", "含在1694.2内", "200W×4(登伟/数魅随机)", "200W", "主光/底子/特写/辅光", "3×2.8m+龙门架(赠)", "无", "无"],
    ]
    write_table(ws3, r, headers3, rows3, [10, 14, 24, 10, 22, 18, 14, 14])

    # ===== Sheet 4: 推流电脑 =====
    ws4 = wb.create_sheet("推流电脑")
    r = write_title(ws4, "推流电脑方案（3个月起租）", "以下为各商家确认或补充的电脑租赁配置")
    headers4 = [
        "商家", "是否含电脑", "电脑类型", "详细配置", "电脑月租(元)",
        "内存", "显卡", "显示器", "含电脑全套折后(元/月)", "备注",
    ]
    rows4 = [
        [
            "租小杨", "是", "台式",
            "i5-12400F / 32G / 512G SSD / RTX4060 / 27寸1K",
            300, "32G", "RTX4060", "27寸1K", 1660,
            "4060显卡；全套1660/月",
        ],
        [
            "昊天", "是", "台式",
            "i5-12400F / 32G / 512G SSD / RTX3060 12G / 27寸1K",
            300, "32G", "RTX3060 12G", "27寸1K", 1770,
            "3个月起租；包邮寄出；全套1770/月",
        ],
        [
            "猫冬", "是", "台式",
            "i5-12400F / 32G / 512G SSD / RTX3060 6G",
            303, "32G", "RTX3060 6G", "未标注", 1997,
            "折后全套1997/月(1694+303)",
        ],
        [
            "洋洋", "是", "台式(单平台)",
            "i5 12代 / 32G / 512G SSD / RTX3060 / 27寸显示器",
            303, "32G", "RTX3060 6G", "27寸", 1997,
            "折后全套1997/月(1694+303)；非天选笔记本",
        ],
        [
            "阿水", "是", "台式",
            "i7 11代 / 32G / 512G SSD / RTX3060 6G",
            479, "32G", "RTX3060 6G", "未标注", 3148.2,
            "含路由+双电源全套",
        ],
        [
            "易希", "否", "-", "需另租(参考303/月同配置台式)", "另租303", "-", "-", "-", "约1997",
            "拍摄端1694.2+另租电脑303",
        ],
    ]
    write_table(ws4, r, headers4, rows4, [10, 10, 12, 32, 12, 8, 14, 10, 16, 28])

    # ===== Sheet 5: 服务支持 =====
    ws5 = wb.create_sheet("服务支持")
    r = write_title(ws5, "服务与支持（3个月租期）", "上门调试与远程协助：六家均已沟通确认提供")
    headers5 = [
        "商家", "最短租期", "上门调试", "日常远程协助", "包邮/物流",
        "非人为免费维修", "押金/满3月退押", "售后时段", "其他说明",
    ]
    rows5 = [
        ["租小杨", "3个月", "是", "是", "需确认", "报价单有", "满3月退押无违约", "8:30-22:30", "不限次调试+维保"],
        ["昊天", "3个月起租", "是", "是", "包邮寄出", "需确认", "需确认", "需确认", "含顶光+魔术腿"],
        ["阿水", "3个月", "是", "是", "单程包邮(报价单)", "需确认", "需确认", "需确认", "折后约7折"],
        ["猫冬", "3个月", "是", "是", "需确认", "需确认", "需确认", "需确认", "折后全套1997/月"],
        ["洋洋", "3个月", "是", "是", "需确认", "需确认", "需确认", "需确认", "折后全套1997/月"],
        ["易希", "3个月", "是", "是", "单程包邮(报价单)", "需确认", "需确认", "需确认", "送绿幕"],
    ]
    write_table(ws5, r, headers5, rows5, [10, 12, 10, 14, 12, 14, 16, 12, 22])

    # ===== Sheet 6: 3个月费用排序 =====
    ws6 = wb.create_sheet("3个月费用排序")
    r = write_title(ws6, "3个月费用排序（折后价）", "按含电脑全套折后月租从低到高排列")
    headers6 = ["排名", "商家", "方案说明", "月租折后(元)", "3个月合计(元)", "适合场景"]
    rows6 = [
        [1, "租小杨", "FX30+Power1000+4060台式32G", 1660, 4980, "含电脑全套最低；户外有电源"],
        [2, "昊天", "A7C+顶光+3060 12G台式32G", 1770, 5310, "性价比高；有顶光；包邮"],
        [3, "猫冬", "A7M4+300B×4+3060台式(折后全套)", 1997, 5991, "A7M4高画质+便宜全套"],
        [4, "洋洋", "A7M4+4灯+3060台式(折后全套)", 1997, 5991, "印迹面光显色好"],
        [5, "易希+另租电脑", "A7M4+绿幕+3060台式", "约1997", "约5991", "室内绿幕为主"],
        [6, "阿水", "FX30+360W灯+路由+双电源+3060", 3148.2, 9444.6, "户外最省心；全套最贵"],
        ["—", "租小杨", "仅拍摄端(折后)", 1360, 4080, "自有电脑"],
        ["—", "阿水", "仅拍摄端(折后)", 1456, 4368, "自有电脑+可另配户外"],
        ["—", "昊天", "仅拍摄端(折后)", 1470, 4410, "自有电脑"],
        ["—", "猫冬/洋洋/易希", "仅拍摄端(折后)", "1694", 5082, "不含电脑"],
    ]
    end = write_table(ws6, r, headers6, rows6, [6, 12, 30, 14, 14, 24])
    for row in range(r + 1, end):
        rank = ws6.cell(row=row, column=1).value
        if rank in (1, 2, 3):
            for col in range(1, 7):
                ws6.cell(row=row, column=col).fill = HIGHLIGHT_FILL

    # ===== Sheet 7: 场景匹配 =====
    ws7 = wb.create_sheet("场景匹配建议")
    r = write_title(
        ws7,
        "场景匹配 — 武隆户外 + 室内绿幕 + 高画质",
        "关键缺口：户外需聚合网络+电源；室内绿幕需幕布+均匀布光",
    )
    headers7 = [
        "商家", "拍摄端折后(元/月)", "含电脑折后(元/月)", "3个月含电脑(元)",
        "户外网络", "户外电源", "绿幕", "画质档", "主要优势", "主要缺口",
    ]
    rows7 = [
        ["租小杨", 1360, 1660, 4980, "无", "Power1000×1", "无", "高(FX30)", "全套最便宜+电源+4060", "无路由/绿幕"],
        ["昊天", 1470, 1770, 5310, "无", "无", "无", "中(A7C)", "顶光+3060 12G+包邮", "无绿幕/户外电源/路由"],
        ["猫冬", 1694, 1997, 5991, "无", "无", "无", "高(A7M4+300B)", "A7M4+便宜全套", "无绿幕/户外配件"],
        ["洋洋", 1694, 1997, 5991, "无", "无", "无", "高(A7M4+印迹)", "面光显色+同价全套", "无绿幕/户外配件"],
        ["易希", 1694.2, "约1997", "约5991", "无", "无", "有(赠)", "高(A7M4)", "唯一送绿幕", "电脑需另租/无户外配件"],
        ["阿水", 1456, 3148.2, 9444.6, "有", "双Power1000", "无", "高(FX30+360W)", "户外全套最完整", "最贵；无绿幕"],
    ]
    write_table(ws7, r, headers7, rows7, [10, 14, 14, 14, 10, 14, 10, 14, 22, 22])

    # ===== Sheet 8: 优缺点 =====
    ws8 = wb.create_sheet("各家优缺点")
    r = write_title(ws8, "各家优缺点（3个月 / 含电脑推流 / 武隆户外+室内绿幕 / 高画质）")
    headers8 = ["商家", "含电脑折后(元/月)", "3个月合计(元)", "优点", "缺点", "综合适合度"]
    rows8 = [
        [
            "租小杨", 1660, 4980,
            "含电脑全套最低价；4060显卡；FX30+Power1000；64G卡+备用电池全；售后条款最强",
            "无绿幕/聚合路由；18-105光圈小；无散热；灯光数量未详",
            "★★★★★ 预算优先+户外电源",
        ],
        [
            "昊天", 1770, 5310,
            "折后1470拍摄端；3060 12G大显存；有顶光+魔术腿；3个月起租包邮寄出；配件全",
            "A7C机身档略低；无绿幕/户外电源/路由；无散热",
            "★★★★ 性价比+顶光高级感",
        ],
        [
            "猫冬", 1997, 5991,
            "折后全套1997；A7M4+VC20Pro；4×300B灯光；32G台式3060；供电配件全",
            "无绿幕/路由/电源；麦克风品牌未写；台式不便携",
            "★★★★ A7M4高画质+稳定全套",
        ],
        [
            "洋洋", 1997, 5991,
            "折后全套1997；A7M4+4灯；印迹200D面光显色好；64G卡；i5 12代+3060+27寸",
            "无绿幕/路由/电源；与猫冬同价但灯光方案不同；麦克风品牌未写",
            "★★★★ 高画质+肤色表现",
        ],
        [
            "易希", "约1997", "约5991",
            "折后1694+送绿幕；A7M4+553采集；灯光分灯位适合抠像；3个月调试包邮",
            "电脑需另租；200W灯品牌随机；无户外网络/电源",
            "★★★★ 室内绿幕比重大",
        ],
        [
            "阿水", 3148.2, 9444.6,
            "唯一含聚合路由+双Power1000；KK360W灯光最强；FX30长播；32G+3060",
            "全套最贵(3个月近9500)；无绿幕；18-50镜头远景弱",
            "★★★★★ 户外为主、预算充足",
        ],
    ]
    write_table(ws8, r, headers8, rows8, [10, 14, 14, 36, 36, 22])

    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_workbook()
    print(f"已生成: {path}")
