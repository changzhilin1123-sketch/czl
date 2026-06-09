# -*- coding: utf-8 -*-
"""Generate ops-list vendor quote comparison Excel."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).resolve().parent / "运营清单七家报价对比.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="微软雅黑", bold=True, size=14, color="1F4E79")
BODY_FONT = Font(name="微软雅黑", size=10)
NOTE_FONT = Font(name="微软雅黑", size=9, color="666666")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
ALT_FILL = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
GREEN_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")


def style_header_row(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def style_data_area(ws, start_row, end_row, col_count, alt=False, highlight_rows=None):
    highlight_rows = highlight_rows or set()
    for r in range(start_row, end_row + 1):
        for c in range(1, col_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c in (1, 5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if r in highlight_rows:
                cell.fill = HIGHLIGHT_FILL
            elif alt and (r - start_row) % 2 == 1:
                cell.fill = ALT_FILL


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_title(ws, title, note=None, merge_cols=12):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=merge_cols)
    row = 2
    if note:
        ws[f"A{row}"] = note
        ws[f"A{row}"].font = NOTE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=merge_cols)
        row += 1
    return row


def write_table(ws, start_row, headers, rows, col_widths, highlight_rows=None):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h)
    style_header_row(ws, start_row, len(headers))
    for i, row_data in enumerate(rows):
        r = start_row + 1 + i
        for c, val in enumerate(row_data, 1):
            ws.cell(row=r, column=c, value=val)
    end_row = start_row + len(rows)
    style_data_area(ws, start_row + 1, end_row, len(headers), alt=True, highlight_rows=highlight_rows)
    set_col_widths(ws, col_widths)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return end_row + 2


def build_workbook():
    wb = Workbook()

    # --- Sheet 1: 总览报价 ---
    ws1 = wb.active
    ws1.title = "总览报价"
    note = (
        "口径：公司运营室外设备清单；折后价；含路由合计已计入。"
        "3个月=月租×3。均未确认推流电脑、户外电源。室内绿幕用自有 Q70。"
        "更新：2026-06-02"
    )
    r = write_title(ws1, "武隆景区 · 运营清单七家报价对比（折后）", note, merge_cols=11)
    headers1 = [
        "排名", "供应商", "套餐说明", "月租(元)", "3个月(元)",
        "采集卡×2", "聚合路由", "300W灯×4", "麦克风", "推荐度", "备注",
    ]
    rows1 = [
        [1, "生机", "A7M5+70-200+四大灯", 3230, 9690, "数魅 VC20", "+380(已计入)", "✓", "猛玛一拖二", "★★★★", "总价最低；路由型号待确认"],
        [2, "洋洋", "A7M5+70-200+四大灯+路由", 3241, 9723, "数联 VC20", "大疆5G版(含)", "✓ 捷宝300W", "猛玛一拖二", "★★★★", "含路由；3个月可退、免费微调"],
        [3, "猫冬·方案一", "A7M5+70-200+四大灯+路由", 3409, 10227, "入门随机", "含", "✓ 凛光300W", "大疆/猛玛随机", "★★★", "采集卡品牌不定"],
        [4, "猫冬·方案二", "方案一+升级采集卡", 3563, 10689, "美乐威 Plus", "含", "✓ 凛光300W", "大疆/猛玛随机", "★★★★★", "综合性价比首选"],
        [5, "昊天", "A7M5+70-200+四大灯", 3680, 11040, "美乐威 Pro", "+440(已计入)", "✓ 300W按200W价", "猛玛一拖二", "★★★★", "Pro采集卡；含绿幕(Q70可忽略)"],
        [6, "阿水", "A7M5+70-200+四大灯+路由", 3890, 11670, "美乐威 Pro", "含", "✓ 300W双色", "猛玛M1充电盒", "★★★", "含散热、绿幕；配置最细"],
        [7, "租小杨", "A7M5+70-200+四大灯+路由", 4072, 12216, "美乐威 Plus", "含", "✓ 凛光300W", "大疆/猛玛随机", "★★", "同档最贵"],
    ]
    highlight = {r + 4 for r, row in enumerate(rows1) if row[1] == "猫冬·方案二"}
    write_table(ws1, r, headers1, rows1, [6, 12, 28, 10, 10, 14, 14, 12, 14, 10, 32], highlight_rows=highlight)

    # --- Sheet 2: 路由加价 ---
    ws2 = wb.create_sheet("路由加价明细")
    r2 = write_title(ws2, "聚合路由器加价对照", "仅生机、昊天为套餐外加价；其余含在月租内。", merge_cols=6)
    headers2 = ["供应商", "套餐折后(元/月)", "路由折后加价(元/月)", "含路由合计(元/月)", "3个月合计(元)", "说明"]
    rows2 = [
        ["生机", 2850, 380, 3230, 9690, "路由单独加租"],
        ["昊天", 3240, 440, 3680, 11040, "路由单独加租"],
        ["洋洋", "—", "含在3241内", 3241, 9723, "大疆多网聚合路由器5G版"],
        ["猫冬·方案一", "—", "含在3409内", 3409, 10227, "多网聚合路由器"],
        ["猫冬·方案二", "—", "含在3563内", 3563, 10689, "多网聚合路由器"],
        ["阿水", "—", "含在3890内", 3890, 11670, "多网聚合路由器495(折后)"],
        ["租小杨", "—", "含在4072内", 4072, 12216, "多网聚合路由器"],
    ]
    write_table(ws2, r2, headers2, rows2, [14, 16, 18, 18, 14, 28])

    # --- Sheet 3: 分项配置 ---
    ws3 = wb.create_sheet("分项配置")
    r3 = write_title(ws3, "各家分项配置明细", merge_cols=10)
    headers3 = [
        "供应商", "机身", "镜头", "采集卡", "路由", "灯光",
        "麦克风", "绿幕", "推流电脑", "户外电源",
    ]
    rows3 = [
        ["生机 3230", "A7M5", "70-200 F2.8 II", "数魅 VC20×2", "聚合路由(+380)", "凛光/捷宝300W×4", "猛玛一拖二", "无", "未含·待确认", "未含"],
        ["洋洋 3241", "A7M5", "70-200 F2.8 II", "数联 VC20×2", "大疆5G聚合(含)", "捷宝300W×4", "猛玛一拖二", "未写", "未含·待确认", "未含"],
        ["猫冬·方案一 3409", "A7M5", "70-200 F2.8 II", "数魅/沣标/品胜随机×2", "多网聚合(含)", "凛光300W×4", "大疆/猛玛随机", "无", "未含·待确认", "未含"],
        ["猫冬·方案二 3563", "A7M5", "70-200 F2.8 II", "美乐威 4K Plus×2", "多网聚合(含)", "凛光300W×4", "大疆/猛玛随机", "无", "未含·待确认", "未含"],
        ["昊天 3680", "A7M5", "70-200 GM II", "美乐威 4K Pro×2", "聚合路由(+440)", "捷宝300W×4(按200W价)", "猛玛一拖二", "送", "未含·待确认", "未含"],
        ["阿水 3890", "A7M5", "70-200 F2.8 II", "美乐威 4K Pro×2", "多网聚合(含)", "300W双色×4", "猛玛M1充电盒", "送3×2.8m", "未含·待确认", "未含"],
        ["租小杨 4072", "A7M5", "70-200 F2.8 II", "美乐威 4K Plus×2", "多网聚合(含)", "凛光300W×4", "大疆/猛玛随机", "无", "未含·待确认", "未含"],
    ]
    write_table(ws3, r3, headers3, rows3, [16, 10, 16, 22, 16, 18, 14, 10, 14, 10])

    # --- Sheet 4: 选型建议 ---
    ws4 = wb.create_sheet("选型建议")
    r4 = write_title(ws4, "选型建议摘要", merge_cols=4)
    headers4 = ["场景", "推荐供应商", "月租(元)", "说明"]
    rows4 = [
        ["预算优先", "生机", 3230, "最便宜；VC20采集卡；无人机RTMP可用"],
        ["便宜+路由明确", "洋洋", 3241, "大疆5G路由；服务条款清楚"],
        ["综合首选", "猫冬·方案二", 3563, "美乐威Plus+路由全含；OBS双路较稳"],
        ["采集卡要好", "昊天 / 阿水", "3680 / 3890", "均为美乐威Pro；昊天比阿水省210/月"],
        ["不建议", "租小杨", 4072, "同配置最贵"],
        ["室内绿幕", "自有 Q70", "—", "不必为绿幕选易希(易希尚未报价)"],
        ["无人机合流", "猫冬② / 昊天 / 阿水", "—", "美乐威档更适合HDMI第二路采集"],
    ]
    end4 = write_table(ws4, r4, headers4, rows4, [22, 16, 12, 40])
    for row_idx in range(r4 + 1, end4 - 1):
        if ws4.cell(row=row_idx, column=1).value == "综合首选":
            for c in range(1, 5):
                ws4.cell(row=row_idx, column=c).fill = GREEN_FILL

    return wb


def main():
    wb = build_workbook()
    wb.save(OUTPUT)
    print(f"已生成: {OUTPUT}")


if __name__ == "__main__":
    main()
